"""Generate synthetic DOCX fixtures, then explicitly opt in to paid model tests."""
import argparse
import hashlib
import io
import json
import tempfile
import time
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]

CASES = {
    'docx-paragraph': dict(title='课程回顾通知', paragraphs=['请于10月17日20:30前提交课程回顾，请大家尽快完成。'],
        deadlines=['2026-10-17T20:30:00+08:00']),
    'docx-table': dict(title='本周课程任务', paragraphs=['请完成下列任务，分别按表格中的时间提交。'],
        table=[['任务','截止日期'],['提交阅读报告','10月20日18点前'],['填写反馈问卷','10月22日前']],
        deadlines=['2026-10-20T18:00:00+08:00','2026-10-22']),
    'docx-image': dict(title='实验总结作业', paragraphs=['请完成并提交实验总结，截止时间见下面的图片。'],
        image=['实验总结提交期限','请于10月19日18:00前提交实验总结。','请大家尽快完成。'],
        deadlines=['2026-10-19T18:00:00+08:00']),
    'docx-partial': dict(title='作业通知', paragraphs=['请于2026年10月21日18点前提交读书报告。'],
        field='本段为尚未解析的动态域缓存，不能作为可靠正文。',
        deadlines=['2026-10-21T18:00:00+08:00'], status='partial_success'),
}


def prepare(folder, font):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from PIL import Image, ImageDraw, ImageFont

    folder.mkdir(parents=True, exist_ok=True)
    receipts = {}
    for name, case in CASES.items():
        document = Document()
        for style_name in ('Normal', 'Title'):
            style = document.styles[style_name]
            style.font.name = 'Microsoft YaHei'; style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.size = Pt(12 if style_name == 'Normal' else 22)
            style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
        document.add_paragraph(case['title'], 'Title')
        for line in case['paragraphs']: document.add_paragraph(line)
        if 'table' in case:
            table = document.add_table(rows=0, cols=2)
            table.style = 'Table Grid'
            for values in case['table']:
                row = table.add_row()
                for cell, value in zip(row.cells, values):
                    cell.text = value
                    borders = OxmlElement('w:tcBorders')
                    for edge in ('top','left','bottom','right'):
                        border = OxmlElement('w:'+edge)
                        for key, val in [('val','single'),('sz','4'),('color','D9D9D9')]: border.set(qn('w:'+key),val)
                        borders.append(border)
                    cell._tc.get_or_add_tcPr().append(borders)
            for cell in table.rows[0].cells:
                shading = OxmlElement('w:shd'); shading.set(qn('w:fill'),'E7EEF5'); cell._tc.get_or_add_tcPr().append(shading)
        if 'image' in case:
            image = Image.new('RGB', (1100, 260), 'white')
            draw = ImageDraw.Draw(image); face = ImageFont.truetype(str(font), 28)
            for i, line in enumerate(case['image']):
                if draw.textlength(line, font=face) > 1020: raise ValueError('Synthetic image text would be clipped')
                draw.text((40, 35+i*48), line, font=face, fill='black')
            buffer = io.BytesIO(); image.save(buffer, format='PNG'); data = buffer.getvalue()
            (folder / (name + '.png')).write_bytes(data)
            document.add_picture(io.BytesIO(data), width=Inches(6))
        if 'field' in case:
            field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'DATE')
            run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = case['field']
            run.append(text); field.append(run); document.add_paragraph()._p.append(field)
        document.core_properties.author = 'Synthetic tests'
        document.core_properties.last_modified_by = 'Synthetic tests'
        path = folder / (name+'.docx'); document.save(path)
        receipts[name] = hashlib.sha256(path.read_bytes()).hexdigest()
    (folder / 'synthetic-receipts.json').write_text(json.dumps(receipts, indent=2), encoding='utf-8')


def acceptable(case, result):
    tasks = result['tasks']
    return (result['status'] == case.get('status','success') and len(tasks) == len(case['deadlines']) and
            sorted(t['deadline']['value'] or '' for t in tasks) == sorted(case['deadlines']) and
            all(t['deadline']['status'] == 'explicit' and t['needs_review'] for t in tasks) and
            (case.get('status') == 'partial_success' and not result['record_results'][0]['coverage']['complete'] or
             case.get('status') is None and all(t['review_reasons'] == ['附件引文及定位需核对原件，请确认。'] for t in tasks)))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument('--prepare', action='store_true'); mode.add_argument('--live', action='store_true')
    parser.add_argument('--input-dir', type=Path, required=True)
    parser.add_argument('--font', type=Path)
    parser.add_argument('--env-file', type=Path, default=ROOT / '.env')
    parser.add_argument('--output', type=Path)
    parser.add_argument('--case', choices=list(CASES))
    args = parser.parse_args()
    if args.prepare:
        if not args.font: parser.error('--prepare requires a local CJK --font')
        prepare(args.input_dir, args.font)
        print('Prepared 4 synthetic DOCX files')
        return 0
    if not args.output: parser.error('--live requires a new output file')
    from ddl_text.codex_client import live_options
    from ddl_text.config import Settings
    from ddl_text.media_pipeline import process_media_request
    from ddl_text.tests.docx_cases import docx_request
    from ddl_text.tests.media_cases import memory_loader
    from ddl_text.tests.run_media_smoke import SyntheticCaptureClient
    receipts = json.loads((args.input_dir/'synthetic-receipts.json').read_text(encoding='utf-8'))
    prepared = []
    for name, case in CASES.items():
        if args.case and name != args.case: continue
        data = (args.input_dir/(name+'.docx')).read_bytes()
        if hashlib.sha256(data).hexdigest() != receipts[name]: parser.error('Synthetic fixture changed; regenerate first')
        prepared.append((name, case, data))
    options = live_options(args.env_file)
    if options['provider'] != 'qiniu' or not options['key'].strip(): parser.error('Verified Qiniu provider required')
    settings = Settings(mode='real', model=options['model'], endpoint=options['endpoint'], timeout_seconds=45)
    results = []
    with args.output.open('x', encoding='utf-8') as output:
        for name, case, data in prepared:
            request = docx_request(data); request['request_id'] = 'synthetic:'+name
            record = request['records'][0]; record.update(input_id='synthetic:'+name, course_or_project='合成测试课程')
            record['attachments'][0]['file_name'] = name+'.docx'
            started = time.monotonic()
            with tempfile.TemporaryDirectory(prefix='niuma-docx-synthetic-') as folder:
                capture = Path(folder)/'draft.json'
                result = process_media_request(request, memory_loader(data), SyntheticCaptureClient(options['key'],str(capture)),settings)
                raw = capture.read_text(encoding='utf-8') if capture.exists() else None
            passed = acceptable(case, result)
            results.append(dict(case=name, synthetic_content=case, docx_sha256=receipts[name], request=request,
                raw_draft=raw, result=result, passed=passed, elapsed_seconds=round(time.monotonic()-started,2)))
            output.seek(0); output.truncate()
            json.dump(dict(tested_at=datetime.now(timezone.utc).isoformat(), model=options['model'],
                scope='Generated DOCX smoke tests only; not private files, recall measurement or WeChat integration.', results=results),
                output, ensure_ascii=False, indent=2); output.flush()
            print(json.dumps(dict(case=name,passed=passed,status=result['status'],deadlines=[t['deadline']['value'] for t in result['tasks']]),ensure_ascii=False),flush=True)
    return 0 if all(r['passed'] for r in results) else 1


if __name__ == '__main__':
    raise SystemExit(main())
